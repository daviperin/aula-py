import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from sklearn.linear_model import LinearRegression
from sklearn.metrics import mean_squared_error, r2_score
from sklearn.model_selection import train_test_split

BASE = r"C:\Users\davif\Documents\Obsidian Vault\FUCAPE\2026-2\Linguagens de Programação\Biblioteca\cc079fef-6c3f-4620-878d-c9a41e69a9f9.xlsx"
OUT = r"C:\Users\davif\Fucape\aula-py\CasoPratico2"
df = pd.read_excel(BASE)

# ---------- gráficos ----------
fig, axs = plt.subplots(1, 2, figsize=(12, 5))
axs[0].scatter(df['Critic_Score'], df['Global_Sales'], alpha=.4)
axs[0].set_title("Nota dos críticos × vendas globais")
axs[0].set_xlabel("Nota dos críticos"); axs[0].set_ylabel("Vendas globais (milhões)")
axs[1].scatter(df['User_Score'], df['Global_Sales'], alpha=.4, color="tab:orange")
axs[1].set_title("Nota dos usuários × vendas globais")
axs[1].set_xlabel("Nota dos usuários"); axs[1].set_ylabel("Vendas globais (milhões)")
plt.tight_layout(); plt.savefig(f"{OUT}\\fig_dispersao.png", dpi=150); plt.close()

base = df.dropna(subset=['Critic_Score', 'User_Score'])
X = base[['Critic_Score', 'User_Score']].values
alvos = [("Global", "Global_Sales"), ("América do Norte", "NA_Sales"), ("Europa", "EU_Sales"), ("Japão", "JP_Sales")]
res, modelos = {}, {}
fig, axs = plt.subplots(2, 2, figsize=(11, 9))
for ax, (nome, col) in zip(axs.ravel(), alvos):
    Xtr, Xte, ytr, yte = train_test_split(X, base[col].values, test_size=0.2, random_state=42)
    m = LinearRegression().fit(Xtr, ytr)
    p = m.predict(Xte)
    res[nome] = (mean_squared_error(yte, p), r2_score(yte, p))
    modelos[nome] = m
    ax.scatter(yte, p, alpha=.4)
    ax.plot([yte.min(), yte.max()], [yte.min(), yte.max()], 'k--')
    ax.set_title(f"{nome} (R² = {res[nome][1]:.2f})")
    ax.set_xlabel("Vendas reais (milhões)"); ax.set_ylabel("Vendas previstas (milhões)")
plt.tight_layout(); plt.savefig(f"{OUT}\\fig_modelos.png", dpi=150); plt.close()

cenarios = [("Jogo A", 90, 8.5), ("Jogo B", 75, 7.0), ("Jogo C", 55, 5.5)]
prev = [modelos["Global"].predict([[c, u]])[0] for _, c, u in cenarios]

# ---------- documento ----------
doc = Document()
doc.styles['Normal'].font.name = 'Calibri'
doc.styles['Normal'].font.size = Pt(11)

def tabela(cab, linhas):
    t = doc.add_table(rows=1, cols=len(cab)); t.style = 'Light Grid Accent 1'
    for i, c in enumerate(cab):
        t.rows[0].cells[i].text = c
    for l in linhas:
        cels = t.add_row().cells
        for i, v in enumerate(l):
            cels[i].text = str(v)
    doc.add_paragraph()

def fig(path, w=6.3):
    doc.add_picture(path, width=Inches(w))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading("Caso Prático: O que explica as vendas de um jogo?", 0)
doc.add_paragraph("Disciplina: Linguagens de Programação\nProfessor: Octavio Locatelli\nAluno: Davi")

doc.add_heading("1. Conhecendo os dados", 1)
doc.add_paragraph(f"A base possui {df.shape[0]} linhas e {df.shape[1]} colunas (um jogo/plataforma por linha).")
doc.add_paragraph("Primeiras observações (colunas principais):")
cols = ['Name', 'Platform', 'Genre', 'NA_Sales', 'EU_Sales', 'JP_Sales', 'Global_Sales', 'Critic_Score', 'User_Score']
h = df[cols].head()
tabela(cols, [[("" if pd.isna(v) else v) for v in r] for r in h.values.tolist()])

doc.add_paragraph("Tipos das variáveis: Name, Platform, Genre, Publisher, Developer e Rating são texto; as demais "
                  "(Year_of_Release, vendas, Critic_Score, Critic_Count, User_Score e User_Count) são numéricas (float64).")
aus = df.isnull().sum()
aus = aus[aus > 0]
doc.add_paragraph("Valores ausentes:")
tabela(["Variável", "Ausentes"], [[k, v] for k, v in aus.items()])
doc.add_paragraph(f"Destaque: Critic_Score tem {aus['Critic_Score']} ausentes e User_Score {aus['User_Score']} "
                  f"(cerca de metade da base), pois muitos jogos não receberam avaliação.")

doc.add_paragraph("Estatísticas descritivas das principais variáveis:")
d = df[['NA_Sales', 'EU_Sales', 'JP_Sales', 'Global_Sales', 'Critic_Score', 'User_Score']].describe().round(2)
tabela(["Estatística"] + list(d.columns), [[i] + list(r) for i, r in zip(d.index, d.values.tolist())])

doc.add_heading("2. Explorando os jogos", 1)
for rot, col in [("maior venda global", "Global_Sales"), ("maior venda na América do Norte", "NA_Sales"),
                 ("maior venda na Europa", "EU_Sales"), ("maior venda no Japão", "JP_Sales")]:
    r = df.loc[df[col].idxmax()]
    doc.add_paragraph(f"{rot[0].upper()+rot[1:]}: {r['Name']} ({r[col]:.2f} milhões de unidades).", style='List Bullet')

doc.add_heading("3. Avaliação × vendas", 1)
fig(f"{OUT}\\fig_dispersao.png")
cc = base['Critic_Score'].corr(base['Global_Sales']); cu = base['User_Score'].corr(base['Global_Sales'])
doc.add_paragraph(
    f"Os gráficos mostram uma nuvem de pontos muito dispersa. Há uma leve tendência de alta nas vendas conforme a nota "
    f"dos críticos sobe (correlação de {cc:.2f}), enquanto a nota dos usuários praticamente não mostra relação "
    f"(correlação de {cu:.2f}). Portanto, a relação existe, mas é fraca, principalmente para a nota dos usuários. "
    f"Muitos jogos com notas altas vendem pouco, e poucos jogos muito vendidos (outliers como Wii Sports) distorcem a escala.")

doc.add_heading("4. Modelo de vendas globais", 1)
doc.add_paragraph(f"Regressão linear com Y = Global_Sales e X = Critic_Score e User_Score, treino 80% / teste 20% "
                  f"(random_state=42). Foram usados apenas os jogos com as duas notas disponíveis ({len(base)} de {len(df)}), pois preencher as notas ausentes com 0 distorceria o modelo.")
doc.add_paragraph(f"Mean Squared Error (MSE): {res['Global'][0]:.2f}", style='List Bullet')
doc.add_paragraph(f"R²: {res['Global'][1]:.2f}", style='List Bullet')

doc.add_heading("5. O modelo consegue prever as vendas?", 1)
doc.add_paragraph("Os gráficos abaixo comparam vendas reais e previstas; a linha tracejada é a previsão perfeita. "
                  "O primeiro painel (Global) responde a esta questão.")
fig(f"{OUT}\\fig_modelos.png", 6.0)
doc.add_paragraph(f"Os pontos estão longe da linha de referência: as previsões ficam concentradas em uma faixa estreita "
                  f"(abaixo de ~1,5 milhão), enquanto as vendas reais variam muito. Com R² de {res['Global'][1]:.2f}, "
                  f"o modelo explica cerca de {res['Global'][1]*100:.0f}% da variação das vendas, ou seja, prevê mal as vendas.")

doc.add_heading("6. O comportamento é igual em diferentes mercados?", 1)
tabela(["Mercado", "MSE", "R²"], [[n, f"{v[0]:.2f}", f"{v[1]:.2f}"] for n, v in res.items()])
doc.add_paragraph(
    f"As avaliações têm baixa capacidade explicativa em todos os mercados; entre os mercados regionais, a maior é na Europa (R² = {res['Europa'][1]:.2f}), seguida da América do Norte ({res['América do Norte'][1]:.2f}); o modelo global tem {res['Global'][1]:.2f}. No Japão o R² é de apenas {res['Japão'][1]:.2f}"
    f", isto é, lá as avaliações quase não ajudam a explicar as vendas. Os MSEs são menores em mercados com vendas menores (Japão, Europa), "
    f"mas isso reflete a escala das vendas, não um modelo melhor; a comparação correta é feita pelo R².")

doc.add_heading("7. Prevendo as vendas de novos jogos", 1)
tabela(["Jogo", "Critic_Score", "User_Score", "Vendas globais previstas (milhões)"],
       [[n, c, str(u).replace('.', ','), f"{p:.2f}"] for (n, c, u), p in zip(cenarios, prev)])
doc.add_paragraph(f"O modelo estima {prev[0]:.2f}, {prev[1]:.2f} e {prev[2]:.2f} milhão de unidades para os jogos A, B e C. "
                  f"O modelo ordena os jogos conforme as notas, mas os valores devem ser lidos com cautela, dado o R² baixo (o modelo erra bastante em jogos individuais).")

doc.add_heading("Conclusão", 1)
doc.add_paragraph(
    f"Não é possível prever o sucesso comercial de um videogame apenas com as avaliações de críticos e usuários. "
    f"A nota dos críticos tem correlação fraca com as vendas ({cc:.2f}) e a dos usuários é quase nula ({cu:.2f}); "
    f"os modelos de regressão apresentaram R² de no máximo {max(v[1] for v in res.values()):.2f}, isto é, explicam apenas cerca de {max(v[1] for v in res.values())*100:.0f}% "
    f"da variação das vendas, e as previsões individuais têm grande margem de erro. "
    f"As vendas dependem de outros fatores, como plataforma, gênero, publisher, marketing e ano de lançamento, que "
    f"deveriam ser incluídos em modelos futuros. Uma limitação adicional é que a análise usa apenas os {len(base)} jogos que têm as duas notas, "
    f"cerca de {len(base)/len(df)*100:.0f}% da base.")

import re
for p in doc.paragraphs:
    if not p.runs: continue
    for r in p.runs: r.text = re.sub(r'(?<=\d)\.(?=\d)', ',', r.text)
doc.save(f"{OUT}\\Relatorio_CasoPratico2_Games.docx")
print("ok", res, prev)
